# Genera las plantillas de Informe de Gestión Institucional 2026 por nivel.
# Origen: análisis de los 14 informes 2025 de "Informes Institucionales".
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x0A, 0x2D, 0x63)
NAVY_DEEP = RGBColor(0x06, 0x1E, 0x46)
BLUE = RGBColor(0x21, 0x66, 0xD6)
TEAL = RGBColor(0x00, 0x96, 0xA5)
GRAY = RGBColor(0x5B, 0x6B, 0x80)

OUT = r"C:\Users\Admin\Documents\Planeación Desarrollos 2026\PDI\Informes Institucionales\Plantillas 2026"
os.makedirs(OUT, exist_ok=True)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)

def p(doc, text="", bold=False, color=None, size=None, align=None, space_after=6, italic=False):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(space_after)
    if align:
        par.alignment = align
    run = par.add_run(text)
    run.bold = bold
    run.italic = italic
    if color: run.font.color.rgb = color
    if size: run.font.size = Pt(size)
    return par

def h(doc, text, level=1):
    sizes = {0: 22, 1: 15, 2: 12.5, 3: 11}
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(14 if level <= 1 else 10)
    par.paragraph_format.space_after = Pt(6)
    run = par.add_run(text)
    run.bold = True
    run.font.size = Pt(sizes.get(level, 11))
    run.font.color.rgb = NAVY if level <= 1 else NAVY_DEEP if level == 2 else BLUE
    if level <= 1:
        pPr = par._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "4"); bottom.set(qn("w:color"), "2166D6")
        bdr.append(bottom); pPr.append(bdr)
    return par

def bullet(doc, text, bold_prefix=None):
    par = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = par.add_run(bold_prefix + " ")
        r.bold = True
    par.add_run(text)
    return par

def auto(doc, text):
    """Bloque de datos que llena el portal automáticamente."""
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.left_indent = Cm(0.3)
    r1 = par.add_run("AUTOMÁTICO · ")
    r1.bold = True; r1.font.color.rgb = TEAL; r1.font.size = Pt(9.5)
    r2 = par.add_run(text)
    r2.font.size = Pt(9.5); r2.font.color.rgb = TEAL
    return par

def aporte(doc, text):
    """Espacio de escritura del encargado."""
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.left_indent = Cm(0.3)
    r1 = par.add_run("APORTE · ")
    r1.bold = True; r1.font.color.rgb = NAVY; r1.font.size = Pt(9.5)
    r2 = par.add_run(text)
    r2.font.size = Pt(9.5); r2.font.color.rgb = GRAY; r2.italic = True
    return par

def guia(doc, text):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(10)
    r = par.add_run(text)
    r.font.size = Pt(9); r.font.color.rgb = GRAY; r.italic = True
    return par

def tabla(doc, headers, rows, widths=None, header_fill="0A2D63"):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, htxt in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        par = cell.paragraphs[0]
        run = par.add_run(htxt)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
        shade(cell, header_fill)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if str(val).startswith("AUTOMÁTICO") or str(val).startswith("[AUTO"):
                run.font.color.rgb = TEAL
    if widths:
        for j, w in enumerate(widths):
            for row in t.rows:
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def portada(doc, nivel, unidad_ejemplo, base_pilares):
    for _ in range(3): doc.add_paragraph()
    p(doc, "UNIVERSIDAD DE SANTANDER", bold=True, color=NAVY_DEEP, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(doc, "INFORME DE GESTIÓN INSTITUCIONAL 2026", bold=True, color=NAVY, size=20, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p(doc, nivel, bold=True, color=TEAL, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    tabla(doc,
          ["Campo", "Valor"],
          [["Unidad", unidad_ejemplo],
           ["Campus / Nivel", "[Bucaramanga · Cúcuta · Valledupar · Institucional]"],
           ["Periodo reportado", "Enero – diciembre de 2026"],
           ["Presentado por", "[Gestor de la unidad — AUTOMÁTICO: organigrama del portal]"],
           ["Fecha de radicación", "[AUTOMÁTICO: fecha de envío en el portal]"]],
          widths=[5, 11])
    doc.add_page_break()
    h(doc, "Cómo se llena esta plantilla", 1)
    p(doc, "Cada sección distingue de dónde viene la información:", bold=True)
    bullet(doc, "lo trae el portal ya calculado: pilares y líneas de acción del PDI, indicadores de la unidad y cifras de «UDES en Cifras». El encargado la revisa y no la edita.", bold_prefix="AUTOMÁTICO (azul-verde) —")
    bullet(doc, "la escribe el encargado de reportar: narrativa, logros, retos, evidencias y anexos.", bold_prefix="APORTE (azul oscuro) —")
    p(doc, "")
    h(doc, "Conexiones con la Plataforma de Planeación", 2)
    tabla(doc,
          ["Conexión", "Qué trae el portal", "Dónde vive"],
          [["1 · PDI", "Pilares estratégicos y líneas de acción en los que trabaja la unidad, con sus proyectos y avance", "Portal de implementación (pilares, líneas de acción, proyectos)"],
           ["2 · Indicadores", "Serie por campus y unidad del catálogo de indicadores (últimos dos años, cumplimiento de meta)", "Tablero de indicadores (v_indicators, v_campus_indicators)"],
           ["3 · UDES en Cifras", "Cifras oficiales del periodo: matrícula, deserción, evaluación profesoral, Saber PRO/TyT, investigación, convenios, rankings, financieras (según la unidad)", "UDES en Cifras (od_indicators / od_values, por campus y periodo)"]],
          widths=[3.2, 7.6, 5.2])
    p(doc, "")
    guia(doc, "Las cifras automáticas se guardan como copia del valor al momento de llenar (foto del periodo): el informe queda estable aunque los datos del portal se actualicen después.")
    doc.add_page_break()

def leyenda_seccion(doc, autos, aportes):
    if autos:
        for a in autos: auto(doc, a)
    if aportes:
        for a in aportes: aporte(doc, a)

# ---------------------------------------------------------------- PLANTILLAS

def plantilla_decanatura():
    doc = Document()
    portada(doc, "Plantilla · Decanatura", "Facultad de [___] (organigrama)", None)

    h(doc, "1. Resumen Ejecutivo", 1)
    guia(doc, "Síntesis de una página: principales logros, retos, cifras clave y hechos relevantes de la gestión anual de la facultad.")
    leyenda_seccion(doc, [], ["Escribir 1 a 3 párrafos. Pilo puede redactar un borrador a partir de los resultados y aportes de las secciones 4 a 7."])

    h(doc, "2. Alineación Estratégica con el PDI", 1)
    leyenda_seccion(doc,
        ["Tabla con los pilares del PDI y las líneas de acción en las que la facultad tuvo proyectos o actividades en 2026, con su estado de avance."],
        ["Narrar cómo las acciones del año contribuyeron a la misión, la visión y esos pilares (media página)."])
    tabla(doc, ["Pilar del PDI", "Línea(s) de acción", "Proyectos / actividades 2026", "Avance"],
          [["AUTOMÁTICO: pilar", "AUTOMÁTICO: línea", "AUTOMÁTICO: proyectos de la facultad", "AUTOMÁTICO: % avance"],
           ["…", "…", "…", "…"]],
          widths=[3.5, 4, 6, 2.5])

    h(doc, "3. Objetivos y Metas 2026 de la Facultad", 1)
    leyenda_seccion(doc, [], ["Objetivo general y metas específicas del año (lista).", "Tabla de metas: valor programado, alcanzado y observación. Si hubo ajustes durante el año, nota que los justifique."])
    tabla(doc, ["Meta", "Valor programado 2026", "Resultado alcanzado", "Observación"],
          [["[meta]", "[valor]", "[valor]", "[cumplida / superada / parcial]"]],
          widths=[6, 3.2, 3.2, 3.6])

    h(doc, "4. Resultados por Dimensión Estratégica", 1)
    guia(doc, "Un bloque por dimensión; en cada uno, primero las cifras que trae el portal y luego los logros y el análisis del encargado.")
    dimensiones = [
        ("4.1 Gestión académica", ["Matrícula por programa y periodo, deserción y retención de la facultad (UDES en Cifras).", "Resultados de aprendizaje: Saber TyT/PRO de los programas (UDES en Cifras)."], ["Logros y análisis del año en la gestión académica."]),
        ("4.2 Gestión profesoral", ["Planta profesoral de la facultad y resultados de la evaluación docente (UDES en Cifras)."], ["Logros, cambios de planta y formación profesoral."]),
        ("4.3 Gestión de la calidad y la oferta", ["Programas con registro calificado vigente y en trámite (si el portal los registra)."], ["Radicaciones, renovaciones y nueva oferta; resultados de auditorías y autoevaluación."]),
        ("4.4 Gestión investigativa", ["Grupos y semilleros categorizados, proyectos y productos (indicadores de investigación)."], ["Logros investigativos de la facultad."]),
        ("4.5 Extensión e internacionalización", ["Convenios, proyectos de extensión y actividades de internacionalización del aula (indicadores de la unidad)."], ["Logros en extensión, convenios y movilidades."]),
        ("4.6 Gestión administrativa y financiera", ["Cifras financieras de la facultad, si aplica (UDES en Cifras)."], ["Resultados administrativos y financieros relevantes."]),
    ]
    for titulo, autos, aportes in dimensiones:
        h(doc, titulo, 2)
        leyenda_seccion(doc, autos, aportes)

    h(doc, "5. Síntesis de Resultados Institucionales", 1)
    auto(doc, "Tabla de indicadores de la facultad con su serie de dos años, meta y cumplimiento (semáforo), agrupada por dimensión estratégica.")
    tabla(doc, ["Indicador", "2025", "2026", "Meta 2026", "Cumplimiento", "Fuente"],
          [["AUTOMÁTICO: indicador", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO: semáforo", "AUTOMÁTICO: sistema"],
           ["…", "…", "…", "…", "…", "…"]],
          widths=[5.2, 1.7, 1.7, 2, 2.4, 3])
    aporte(doc, "Comentario de la decanatura sobre la síntesis (medio página).")

    h(doc, "6. Puntos Críticos y Acciones de Mejora", 1)
    leyenda_seccion(doc, [], ["Tabla: punto crítico, acción propuesta, tiempo estimado y recursos necesarios."])
    tabla(doc, ["Dimensión estratégica", "Punto crítico", "Acción de mejora", "Tiempo", "Recursos"],
          [["[dimensión]", "[punto crítico]", "[acción]", "[2026 / 2027]", "[recursos]"]],
          widths=[3.4, 4, 4.4, 2, 2.2])

    h(doc, "7. Logros Destacados", 1)
    guia(doc, "Máximo 3 logros extraordinarios del año (innovaciones, reconocimientos, proyectos disruptivos), cada uno con contexto, resultado y evidencia.")

    h(doc, "8. Evidencias y Fuentes de Verificación", 1)
    leyenda_seccion(doc, [], ["Tabla: logro o resultado, fuente de verificación, anexo o enlace. Los anexos se cargan en el portal y se numeran solos."])
    tabla(doc, ["Logro / Resultado", "Fuente de verificación", "Anexo (portal)"],
          [["[logro]", "[fuente]", "AUTOMÁTICO: enlace al anexo cargado"]],
          widths=[5.5, 5.5, 5])

    h(doc, "9. Conclusiones", 1)
    leyenda_seccion(doc, [], ["1 a 3 párrafos: valoración del año, aporte a la misión y rumbo del próximo periodo."])

    h(doc, "10. Anexos", 1)
    auto(doc, "Lista numerada de los anexes cargados en el portal, con su autor y fecha; cada uno queda citado desde el cuerpo del informe.")
    return doc

def plantilla_vicerrectoria():
    doc = Document()
    portada(doc, "Plantilla · Vicerrectoría u Oficina", "Vicerrectoría de [___] (organigrama)", None)

    h(doc, "1. Resumen Ejecutivo", 1)
    guia(doc, "Síntesis de una página (máximo dos): logros, retos, cifras clave, oportunidades y hechos relevantes de la gestión anual.")
    leyenda_seccion(doc, [], ["1 a 3 párrafos. Pilo puede redactar un borrador con base en las secciones 3 a 7."])

    h(doc, "2. Alineación Estratégica", 1)
    guia(doc, "Cómo las acciones y resultados del año contribuyeron a la misión, la visión y los pilares del PDI. Media página; viñetas si hay varios pilares.")
    leyenda_seccion(doc,
        ["Pilares del PDI con líneas de acción en las que la unidad tuvo proyectos en 2026 y su avance."],
        ["Narrativa de alineación estratégica."])
    tabla(doc, ["Pilar del PDI", "Línea(s) de acción", "Proyectos 2026", "Avance"],
          [["AUTOMÁTICO: pilar", "AUTOMÁTICO: línea", "AUTOMÁTICO: proyectos", "AUTOMÁTICO: % avance"]],
          widths=[3.5, 4.5, 5.5, 2.5])

    h(doc, "3. Objetivos y Metas Anuales de la Unidad", 1)
    leyenda_seccion(doc, [], ["Objetivo general y metas específicas 2026 (lista).", "Tabla de metas alcanzadas; nota justificando ajustes si los hubo."])
    tabla(doc, ["Meta", "Programado 2026", "Alcanzado", "Observación"],
          [["[meta]", "[valor]", "[valor]", "[cumplida / superada / parcial]"]],
          widths=[6, 3.2, 3.2, 3.6])

    h(doc, "4. Resultados y Evidencias", 1)
    guia(doc, "Logros y actividades del año con datos y evidencias verificables; comparar con el año previo (1 a 2 páginas).")
    leyenda_seccion(doc,
        ["Matriz de resultados: indicadores de la unidad con serie 2025–2026, variación y fuente (catálogo de indicadores del portal).",
         "Cifras oficiales del tema de la unidad en UDES en Cifras (p. ej. publicaciones y grupos para investigación; convenios y emprendimientos para extensión; cifras NIIF para la VAF)."],
        ["Resultados principales (viñetas).", "Logros extraordinarios y narrativa de los resultados."])
    tabla(doc, ["Indicador", "2025", "2026", "Variación", "Fuente de verificación"],
          [["AUTOMÁTICO: indicador", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO: sistema o anexo"],
           ["…", "…", "…", "…", "…"]],
          widths=[5.4, 2, 2, 2.4, 4.2])
    h(doc, "Resultados financieros (solo VAF)", 3)
    aporte(doc, "Estado de situación financiera, endeudamiento y excedentes, con notas de corte y periodo. Las cifras base puede traerlas UDES en Cifras.")

    h(doc, "5. Logros Innovadores o Excepcionales", 1)
    guia(doc, "Máximo 3: innovaciones, premios, reconocimientos o proyectos disruptivos, con contexto, resultado y evidencia documental.")

    h(doc, "6. Análisis de Retos y Oportunidades", 1)
    guia(doc, "Media a una página. Se recomienda citar colaboraciones interáreas y aprendizajes compartidos.")
    leyenda_seccion(doc, [], ["Retos enfrentados (lista).", "Aprendizajes obtenidos (lista).", "Oportunidades detectadas (lista).", "Propuestas para el próximo periodo (lista)."])

    h(doc, "7. Evidencia y Fuentes de Verificación", 1)
    leyenda_seccion(doc, [], ["Tabla: logro/indicador, fuente de verificación, anexo o enlace y observaciones."])
    tabla(doc, ["Logro / Indicador", "Fuente de verificación", "Anexo (portal)", "Observaciones"],
          [["[logro]", "[fuente]", "AUTOMÁTICO: enlace al anexo", "[observación]"]],
          widths=[4.6, 4.4, 3.5, 3.5])

    h(doc, "8. Conclusiones", 1)
    guia(doc, "1 a 3 párrafos, con enfoque claro, positivo y propositivo.")

    h(doc, "9. Anexos", 1)
    auto(doc, "Lista numerada de los anexos cargados en el portal; cada anexo queda referenciado desde el cuerpo del informe.")
    return doc

def plantilla_rectoria_campus():
    doc = Document()
    portada(doc, "Plantilla · Rectoría de Campus", "Rectoría Campus [Bucaramanga / Cúcuta / Valledupar]", None)

    h(doc, "1. Resumen Ejecutivo del Campus", 1)
    guia(doc, "Síntesis global del año del campus: logros, retos y cifras clave de todas sus unidades.")
    leyenda_seccion(doc,
        ["Cifras generales del campus (matrícula, deserción, graduados, investigación, convenios) del periodo, desde UDES en Cifras."],
        ["1 a 3 párrafos del rector del campus. Pilo puede redactar el borrador con base en los informes recibidos."])

    h(doc, "2. Alineación Estratégica con el PDI", 1)
    leyenda_seccion(doc,
        ["Pilares y líneas de acción del PDI con avance de los proyectos del campus."],
        ["Narrativa de alineación del campus con misión, visión y PDI."])

    h(doc, "3. Objetivos y Metas del Campus", 1)
    leyenda_seccion(doc, [], ["Objetivo general y metas 2026 del campus.", "Tabla de metas alcanzadas."])
    tabla(doc, ["Meta", "Programado 2026", "Alcanzado", "Observación"],
          [["[meta]", "[valor]", "[valor]", "[cumplida / superada / parcial]"]],
          widths=[6, 3.2, 3.2, 3.6])

    h(doc, "4. Consolidado por Función Sustantiva", 1)
    guia(doc, "Un bloque por función; cada uno reúne lo enviado por las unidades del campus (el portal lo arma con los informes ya aprobados) y deja espacio para el análisis del rector.")
    funciones = [
        ("4.1 Investigación", "Grupos categorizados, proyectos, productos y_publicaciones del campus (UDES en Cifras + indicadores)."),
        ("4.2 Extensión", "Convenios, emprendimientos, proyectos de extensión y movilidades del campus."),
        ("4.3 Enseñanza", "Matrícula, deserción, evaluación profesoral y Saber PRO/TyT del campus."),
        ("4.4 Gestión Administrativa y Financiera", "Cifras operativas y financieras del campus (si aplican)."),
    ]
    for titulo, autos in funciones:
        h(doc, titulo, 2)
        auto(doc, "Aportes aprobados de las unidades del campus en esta función (resumen por unidad).")
        auto(doc, autos)
        aporte(doc, "Análisis y destacados del rector del campus sobre esta función.")

    h(doc, "5. Resultados y Evidencias", 1)
    auto(doc, "Matriz de indicadores del campus con serie 2025–2026, variación y semáforo de cumplimiento.")
    tabla(doc, ["Indicador", "2025", "2026", "Meta 2026", "Cumplimiento", "Fuente"],
          [["AUTOMÁTICO: indicador", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO", "AUTOMÁTICO: semáforo", "AUTOMÁTICO: sistema"],
           ["…", "…", "…", "…", "…", "…"]],
          widths=[5.2, 1.7, 1.7, 2, 2.4, 3])
    aporte(doc, "Logros extraordinarios del campus (máximo 3), con evidencia.")

    h(doc, "6. Análisis de Retos y Oportunidades", 1)
    leyenda_seccion(doc, [], ["Retos enfrentados, aprendizajes, oportunidades y propuestas para el próximo periodo."])

    h(doc, "7. Conclusiones", 1)
    leyenda_seccion(doc, [], ["1 a 3 párrafos."])

    h(doc, "8. Anexos", 1)
    auto(doc, "Lista numerada de anexos cargados; además, enlaces a los informes aprobados de las unidades del campus (decanaturas, vicerrectorías y oficinas).")
    return doc

def plantilla_rectoria_general():
    doc = Document()
    portada(doc, "Plantilla · Rectoría General", "Rectoría General UDES (nivel institucional)", None)

    h(doc, "Presentación del Rector", 1)
    guia(doc, "Mensaje del Rector General: sentido del año, mandatos institucionales y resultados más relevantes (1 página).")
    leyenda_seccion(doc, [], ["Escrito por el Rector General. Pilo puede proponer un borrador a partir de la síntesis de los informes de campus y vicerrectorías."])

    h(doc, "Introducción", 1)
    leyenda_seccion(doc, [], ["Propósito del informe, periodo cubierto y organización del documento (2 a 3 párrafos)."])

    for num, fn, autos in [
        ("I", "Investigación", ["Resultados por pilar institucional: grupos categorizados, publicaciones Scopus, proyectos, financiación (UDES en Cifras + indicadores de investigación).", "Aportes aprobados de la Vicerrectoría/Consejo de Investigaciones y de los campus."]),
        ("II", "Extensión", ["Resultados por pilar: convenios, emprendimientos, proyectos con el territorio y movilidades.", "Aportes aprobados de la Vicerrectoría de Extensión y de los campus."]),
        ("III", "Enseñanza", ["Resultados por pilar: matrícula, deserción, evaluación profesoral, Saber PRO/TyT, nueva oferta y registros calificados.", "Aportes aprobados de la Vicerrectoría de Enseñanza y de los campus."]),
        ("IV", "Gestión Administrativa y Financiera", ["Cifras NIIF del ejercicio: activos, pasivos, ingresos, excedentes e índice de endeudamiento (UDES en Cifras).", "Aportes aprobados de la VAF."]),
    ]:
        h(doc, f"{num}. {fn.upper()}", 1)
        auto(doc, "Tabla de «Resultados por pilares institucionales»: un bloque por pilar del PDI con los logros y cifras de la función.")
        for a in autos: auto(doc, a)
        h(doc, "Retos y Oportunidades", 2)
        aporte(doc, "Retos identificados, aprendizajes y oportunidades estratégicas de la función (listas).")
        auto(doc, "Infografía de cifras de la función (la genera el portal desde UDES en Cifras, como hoy en los boletines).")

    h(doc, "V. Síntesis y Proyecciones Institucionales 2027", 1)
    auto(doc, "Semáforo institucional: cumplimiento de los indicadores del PDI por pilar, consolidado de todos los campus y unidades.")
    aporte(doc, "Balance del año y proyecciones estratégicas del próximo periodo (1 a 2 páginas).")

    h(doc, "Anexos", 1)
    auto(doc, "Lista de anexos cargados y enlaces a los informes aprobados de campus, vicerrectorías, oficinas y decanaturas.")
    return doc

PLANTILLAS = [
    ("Plantilla A - Decanatura.docx", plantilla_decanatura),
    ("Plantilla B - Vicerrectoria u Oficina.docx", plantilla_vicerrectoria),
    ("Plantilla C - Rectoria de Campus.docx", plantilla_rectoria_campus),
    ("Plantilla D - Rectoria General.docx", plantilla_rectoria_general),
]

for nombre, fn in PLANTILLAS:
    doc = fn()
    path = os.path.join(OUT, nombre)
    doc.save(path)
    print("OK", nombre)
